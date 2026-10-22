import streamlit as st
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import io

# 1. إعدادات الصفحة الأساسية
st.set_page_config(
    page_title="منصة حصين | حماية وتعمية الوثائق القانونية",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# 2. التصميم الفاخر (كحلي وذهبي) ودعم المحاذاة العربية (RTL)
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@400;500;700;800&display=swap');
    
    html, body, [class*="css"] {
        font-family: 'Tajawal', sans-serif;
        direction: rtl;
        text-align: right;
    }
    
    /* ترويسة المنصة الفاخرة */
    .main-header {
        background: linear-gradient(135deg, #0d1b2a 0%, #1b263b 100%);
        padding: 35px;
        border-radius: 12px;
        border-bottom: 4px solid #c5a059;
        color: #ffffff;
        text-align: center;
        margin-bottom: 25px;
        box-shadow: 0 6px 15px rgba(0,0,0,0.15);
    }
    .main-header h1 {
        color: #c5a059;
        font-weight: 800;
        font-size: 2.6rem;
        margin-bottom: 10px;
    }
    .main-header p {
        color: #e0e1dd;
        font-size: 1.15rem;
        margin: 0;
    }
    
    /* الصناديق والبطاقات */
    .legal-box {
        background-color: #f8f9fa;
        border-right: 5px solid #c5a059;
        padding: 20px;
        border-radius: 8px;
        margin-bottom: 20px;
        box-shadow: 0 2px 5px rgba(0,0,0,0.05);
    }
    .disclaimer-box {
        background-color: #fff8e6;
        border: 1px solid #f0d78c;
        border-right: 5px solid #d97706;
        padding: 15px;
        border-radius: 8px;
        font-size: 0.95rem;
        color: #333;
        margin-bottom: 20px;
    }
    
    /* تحسين شكل الأزرار */
    .stButton>button, .stDownloadButton>button {
        background-color: #0d1b2a !important;
        color: #c5a059 !important;
        border: 1px solid #c5a059 !important;
        font-weight: bold !important;
        border-radius: 8px !important;
        padding: 10px 24px !important;
        width: 100%;
        transition: 0.3s;
    }
    .stButton>button:hover, .stDownloadButton>button:hover {
        background-color: #c5a059 !important;
        color: #0d1b2a !important;
    }
</style>
""", unsafe_allow_html=True)

# 3. عرض الترويسة (Header)
st.markdown("""
<div class="main-header">
    <h1>⚖️ منصة حصين (Haseen)</h1>
    <p>النظام الآلي المتقدم لتعمية وحماية البيانات الحساسة في الوثائق والصكوك القضائية</p>
</div>
""", unsafe_allow_html=True)

# 4. القائمة الجانبية (إعدادات التعمية وسياسة الخصوصية)
with st.sidebar:
    st.markdown("### ⚙️ إعدادات الفحص والتعمية")
    st.write("حدد نوع البيانات الحساسة المراد إخفاؤها من الوثيقة:")
    
    mask_id = st.checkbox("أرقام الهويات والإقامات والسجلات (10 أرقام)", value=True)
    mask_deeds = st.checkbox("أرقام الصكوك والقرارات القضائية", value=True)
    mask_dates = st.checkbox("التواريخ (الهجرية والميلادية)", value=True)
    mask_phones = st.checkbox("أرقام الهواتف والجوالات", value=True)
    mask_iban = st.checkbox("الحسابات البنكية (IBAN)", value=True)
    
    st.markdown("---")
    st.markdown("### ✍️ تعمية أسماء أو عبارات محددة")
    st.caption("أدخل أسماء الأطراف، الشركات، أو أي بيانات ترغب في حجبها تماماً (ضع كل اسم في سطر مستقل):")
    custom_names = st.text_area("قائمة الأسماء المراد إخفاؤها:", placeholder="مثال:\nسعد بن عبدالله\nشركة الأفق المحدودة")
    
    st.markdown("---")
    st.markdown("### 🛡️ ميثاق الأمان والخصوصية")
    st.info("""
    **معالجة فورية وآمنة:**
    تعمل منصة **حصين** بنظام المعالجة اللحظية في الذاكرة العشوائية المؤقتة (RAM). 
    لا يتم حفظ أو تخزين أو أرشفة أي مستندات أو نصوص يتم رفعها على الخادم إطلاقاً، وتُحذف كافة البيانات تلقائياً بمجرد إغلاق الصفحة.
    """)

# 5. شرح عمل المنصة للمحامي
st.markdown("""
<div class="legal-box">
    <h4>مرحباً بك في منصة حصين لحماية الوثائق القانونية</h4>
    <p>صُممت هذه الأداة خصيصاً لمكاتب المحاماة والإدارات القانونية لتوفير بيئة آمنة لتنقيح المستندات واللوائح قبل مشاركتها مع الأطراف الخارجية أو رفعها لأنظمة الذكاء الاصطناعي لدراستها وتحليلها. تقوم الأداة بمسح المستند واستبدال البيانات الحساسة برموز محجوبة <b>[محجوب]</b> مع الحفاظ على سياق النص القانوني.</p>
</div>
""", unsafe_allow_html=True)

# 6. سياسة الاستخدام وإخلاء المسؤولية الإلزامي
st.markdown("""
<div class="disclaimer-box">
    <h4>⚠️ سياسة الاستخدام وإخلاء المسؤولية القانونية:</h4>
    <ul>
        <li><b>طبيعة الأداة:</b> منصة "حصين" هي أداة تقنية مساندة تعتمد على الخوارزميات والأنماط البرمجية القياسية لرصد وتعمية البيانات الحساسة.</li>
        <li><b>المراجعة البشرية:</b> نظراً لاختلاف صياغات الصكوك والوثائق أو احتمالية وجود أخطاء إملائية في المستند الأصلي، لا تضمن المنصة إخفاء البيانات بنسبة 100% بشكل آلي مطلق، وتقع المسؤولية المهنية والقانونية الكاملة على عاتق المستخدم لمراجعة المخرجات النهائية والتحقق منها قبل تداولها أو نشرها.</li>
        <li><b>حدود المسؤولية:</b> يُخلي مطورو المنصة وإدارتها مسؤوليتهم القانونية عن أي أضرار مباشرة أو غير مباشرة قد تنتج عن الاعتماد الكلي على المخرجات الآلية دون تدقيق ومراجعة نهائية من قِبل المختص القانوني.</li>
    </ul>
</div>
""", unsafe_allow_html=True)

agree_terms = st.checkbox("أقر بأنني قرأت وفهمت سياسة الاستخدام وإخلاء المسؤولية، وأوافق على مراجعة المستند المعالج قبل اعتماده رسمياً.")

# 7. دوال المعالجة والتعمية والملفات
def redact_sensitive_data(text, redactions_list):
    redacted_text = text
    
    # 1. تعمية الأسماء المخصصة المدخلة من قِبل المحامي
    if custom_names:
        names = [n.strip() for n in custom_names.split('\n') if n.strip()]
        for name in names:
            if name in redacted_text:
                redacted_text = redacted_text.replace(name, "[اسم طرف محجوب]")
                redactions_list.append(name)
                
    # 2. تعمية رقم الآيبان السعودي (SA + 22 خانة)
    if mask_iban:
        iban_pattern = r'SA\d{22}\b'
        matches = re.findall(iban_pattern, redacted_text)
        redactions_list.extend(matches)
        redacted_text = re.sub(iban_pattern, '[حساب آيبان محجوب]', redacted_text)

    # 3. تعمية الهويات الوطنية والإقامات والسجلات التجارية (10 أرقام تبدأ بـ 1 أو 2 أو 7)
    if mask_id:
        id_pattern = r'\b[127]\d{9}\b'
        matches = re.findall(id_pattern, redacted_text)
        redactions_list.extend(matches)
        redacted_text = re.sub(id_pattern, '[رقم هوية/سجل محجوب]', redacted_text)
        
    # 4. تعمية أرقام الجوالات السعودية
    if mask_phones:
        phone_pattern = r'\b(?:05\d{8}|9665\d{8}|\+9665\d{8})\b'
        matches = re.findall(phone_pattern, redacted_text)
        redactions_list.extend(matches)
        redacted_text = re.sub(phone_pattern, '[رقم هاتف محجوب]', redacted_text)

    # 5. تعمية التواريخ الهجرية والميلادية (مثل 1445/05/12 أو 2024-01-15)
    if mask_dates:
        date_pattern = r'\b(?:\d{4}[/-]\d{1,2}[/-]\d{1,2}|\d{1,2}[/-]\d{1,2}[/-]\d{4})\b'
        matches = re.findall(date_pattern, redacted_text)
        redactions_list.extend(matches)
        redacted_text = re.sub(date_pattern, '[تاريخ محجوب]', redacted_text)

    # 6. تعمية أرقام الصكوك الطويلة (من 11 إلى 14 رقماً)
    if mask_deeds:
        deed_pattern = r'\b\d{11,14}\b'
        matches = re.findall(deed_pattern, redacted_text)
        redactions_list.extend(matches)
        redacted_text = re.sub(deed_pattern, '[رقم صك/وثيقة محجوب]', redacted_text)
        
    return redacted_text

def generate_word_docx(processed_text):
    doc = Document()
    
    # إضافة ترويسة رسمية داخل ملف الوورد
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_para.add_run("وثيقة معالجة ومحمية بواسطة منصة حصين (Haseen)")
    header_run.bold = True
    header_run.font.size = Pt(14)
    header_run.font.color.rgb = RGBColor(197, 160, 89)
    
    doc.add_paragraph("─" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # إضافة النص المعالج مع دعم المحاذاة لليمين
    for line in processed_text.split('\n'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(line)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def redact_pdf_directly(pdf_bytes, items_to_redact):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    unique_items = set(items_to_redact)
    
    for page in doc:
        for item in unique_items:
            if item.strip():
                # البحث عن مكان النص الحساس في صفحة الـ PDF وطمسه باللون الأسود
                areas = page.search_for(item)
                for area in areas:
                    page.add_redact_annot(area, fill=(0, 0, 0))
        page.apply_redactions()
        
    pdf_buffer = io.BytesIO()
    doc.save(pdf_buffer)
    pdf_buffer.seek(0)
    return pdf_buffer

# 8. قسم رفع ومعالجة الملفات (يظهر فقط بعد الموافقة على الإقرار)
if agree_terms:
    st.markdown("### 📂 رفع الوثيقة القانونية (PDF)")
    uploaded_file = st.file_uploader("اختر ملف الـ PDF المراد تعمية بياناته:", type=["pdf"])
    
    if uploaded_file is not None:
        pdf_bytes = uploaded_file.read()
        
        with st.spinner("جاري فحص المستند وتطبيق خوارزميات التعمية..."):
            # قراءة النصوص من ملف الـ PDF
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            full_original_text = ""
            for page in doc:
                full_original_text += page.get_text("text") + "\n"
                
            # تطبيق التعمية ورصد العناصر المحجوبة
            redacted_items = []
            processed_text = redact_sensitive_data(full_original_text, redacted_items)
            
        st.success(f"✅ تمت معالجة الوثيقة بنجاح! تم رصد وحجب ({len(redacted_items)}) عنصر حساس.")
        
        # عرض النص قبل وبعد للمراجعة السريعة
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("#### 📄 النص الأصلي المستخرج")
            st.text_area("النص قبل المعالجة:", value=full_original_text, height=350, disabled=True)
            
        with col2:
            st.markdown("#### 🛡️ النص المحمي (يمكنك التعديل عليه يدوياً هنا قبل التحميل)")
            final_edited_text = st.text_area("النص بعد التعمية:", value=processed_text, height=350)
            
        st.markdown("---")
        st.markdown("### 📥 تصدير وتحميل الوثيقة المحمية")
        
        export_col1, export_col2 = st.columns(2)
        
        with export_col1:
            # تجهيز ملف الوورد
            word_file = generate_word_docx(final_edited_text)
            st.download_button(
                label="📄 تحميل كملف Word (.docx) قابل للتعديل والطباعة",
                data=word_file,
                file_name=f"Haseen_Protected_{uploaded_file.name.replace('.pdf', '')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
        with export_col2:
            # تجهيز ملف الـ PDF المطموس
            redacted_pdf_file = redact_pdf_directly(pdf_bytes, redacted_items)
            st.download_button(
                label="📕 تحميل كملف PDF (طمس أسود مباشر على الأرقام والأسماء)",
                data=redacted_pdf_file,
                file_name=f"Haseen_Redacted_{uploaded_file.name}",
                mime="application/pdf"
            )
else:
    st.warning("يرجى وضع علامة صح (☑️) للموافقة على سياسة الاستخدام وإخلاء المسؤولية أعلاه لتفعيل نظام رفع ومعالجة الوثائق.")

# 9. تذييل الصفحة (Footer)
st.markdown("---")
st.markdown("""
<div style="text-align: center; color: #666; font-size: 0.9rem; padding: 10px;">
    جميع الحقوق محفوظة © منصة حصين (Haseen) لحلول التقنية القانونية | تعمل الأداة بنظام الأمان اللحظي ولا تحتفظ بأي نسخ من الوثائق.
</div>
""", unsafe_allow_html=True)
